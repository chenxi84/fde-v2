"""平台内置文件工具（仿 v1，适配 v2 约定）。

为每个 FDE 应用统一提供约定资源目录与文件访问能力，支撑"用 Agent 做数据导入解析"：
- `resource/import-file/`  人工上传的待导入文件（只读：仅页面上传，工具不可写）
- `resource/export-file/`  Agent / 服务产出的结果文件（工具可写）

对外工具名 = `<应用名>__platform_list_files / __platform_read_file / __platform_write_file`，
Agent 与 MCP 走同一套；执行经 `call_builtin(app_dir, service, params)`，
遵循 v2 返回约定（成功返回业务值 dict，失败抛 FdeError）。
"""
from pathlib import Path

from fde import FdeError

# ── 常量 ────────────────────────────────────────────────

ALLOWED_DIRS = ("import-file", "export-file")
IMPORT_DIR = "import-file"
EXPORT_DIR = "export-file"

TEXT_SUFFIXES = {".csv", ".tsv", ".txt", ".json", ".md", ".log", ".xml", ".yaml", ".yml", ""}
MAX_TEXT_CHARS = 20000      # 文本截断上限（字符）
MAX_XLSX_ROWS = 200         # Excel 截断上限（行）
MAX_WRITE_CHARS = 1_000_000  # 写文件内容上限（字符）

# 平台保留的内置服务名（对外拼接为 <应用名>__<服务名>）
BUILTIN_SERVICES = ("platform_list_files", "platform_read_file", "platform_write_file")


def is_builtin_service(service: str) -> bool:
    return service in BUILTIN_SERVICES


# ── 目录与路径安全 ──────────────────────────────────────


def ensure_resource_dirs(app_dir: Path) -> None:
    """平台加载应用时调用：自动创建约定资源目录（幂等）。"""
    for d in ALLOWED_DIRS:
        (Path(app_dir) / "resource" / d).mkdir(parents=True, exist_ok=True)


def safe_filename(name: str) -> str:
    """清理文件名：保留中文，剥离路径分隔符 / .. / 控制字符，防路径穿越。"""
    name = (name or "").replace("\\", "/").split("/")[-1]
    name = "".join(ch for ch in name if ord(ch) >= 32)
    name = name.replace("..", "").strip().strip(". ")
    return name


def _resolve(app_dir: Path, directory: str, file_name: str, must_exist: bool = True) -> Path:
    """校验目录与文件名，解析出 resource/{directory}/ 下的目标路径；非法即抛 FdeError。"""
    if directory not in ALLOWED_DIRS:
        raise FdeError(f"directory 必须是 {list(ALLOWED_DIRS)} 之一")
    name = safe_filename(file_name)
    if not name:
        raise FdeError(f"非法文件名：{file_name!r}")
    base = (Path(app_dir) / "resource" / directory).resolve()
    target = (base / name).resolve()
    if target.parent != base:  # 路径穿越防护：必须直接位于约定目录下
        raise FdeError(f"非法文件路径：{file_name!r}")
    if must_exist and not target.is_file():
        raise FdeError(f"文件不存在：{directory}/{name}（可先 platform_list_files 查看）")
    return target


# ── 执行器（v2 返回风格）────────────────────────────────


def list_files(app_dir: Path, directory: str = IMPORT_DIR) -> dict:
    if directory not in ALLOWED_DIRS:
        raise FdeError(f"directory 必须是 {list(ALLOWED_DIRS)} 之一")
    target_dir = Path(app_dir) / "resource" / directory
    files = []
    if target_dir.exists():
        for f in sorted(target_dir.iterdir()):
            if f.is_file() and not f.name.startswith("."):
                st = f.stat()
                files.append({"name": f.name, "size": st.st_size, "mtime": st.st_mtime})
    return {"directory": directory, "count": len(files), "files": files}


def read_file(app_dir: Path, directory: str, file_name: str) -> dict:
    target = _resolve(app_dir, directory, file_name)
    size = target.stat().st_size
    suffix = target.suffix.lower()

    if suffix in TEXT_SUFFIXES:
        raw = target.read_text(encoding="utf-8", errors="replace")
        truncated = len(raw) > MAX_TEXT_CHARS
        return {
            "name": target.name, "size": size, "truncated": truncated,
            "content": raw[:MAX_TEXT_CHARS],
        }

    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise FdeError("openpyxl 未安装，无法解析 xlsx（pip install openpyxl）")
        wb = load_workbook(str(target), read_only=True, data_only=True)
        sheets = list(wb.sheetnames)
        ws = wb[sheets[0]] if sheets else None
        rows = []
        if ws is not None:
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_XLSX_ROWS:
                    break
                rows.append("\t".join("" if c is None else str(c) for c in row))
        total = (ws.max_row if ws is not None else 0) or len(rows)
        wb.close()
        return {
            "name": target.name, "size": size, "sheets": sheets,
            "truncated": total > MAX_XLSX_ROWS, "content": "\n".join(rows),
        }

    raise FdeError(f"不支持的文件格式：{suffix or '（无扩展名）'}（二进制请交给应用服务处理）")


def write_file(app_dir: Path, file_name: str, content: str, append: bool = False) -> dict:
    """仅允许写入 export-file（import-file 是人工上传区）。"""
    if not isinstance(content, str):
        raise FdeError("content 必须是字符串（文本内容）")
    if len(content) > MAX_WRITE_CHARS:
        raise FdeError(f"内容超过上限（{MAX_WRITE_CHARS} 字符），请拆分或改用导出服务")
    target = _resolve(app_dir, EXPORT_DIR, file_name, must_exist=False)
    target.parent.mkdir(parents=True, exist_ok=True)
    with open(target, "a" if append else "w", encoding="utf-8", newline="") as f:
        f.write(content)
    return {
        "name": target.name, "directory": EXPORT_DIR, "size": target.stat().st_size,
        "message": f"已{'追加' if append else '写入'} resource/{EXPORT_DIR}/{target.name}",
    }


def call_builtin(app_dir: Path, service: str, params: dict):
    """内置工具执行分发。"""
    params = params or {}
    if service == "platform_list_files":
        return list_files(app_dir, params.get("directory", IMPORT_DIR))
    if service == "platform_read_file":
        return read_file(app_dir, params.get("directory", ""), params.get("file_name", ""))
    if service == "platform_write_file":
        return write_file(
            app_dir, params.get("file_name", ""), params.get("content", ""),
            bool(params.get("append", False)),
        )
    raise FdeError(f"未知内置工具：{service}")


# ── 对外工具定义（Agent 与 MCP 共用）────────────────────


def builtin_tool_defs(prefix: str, *, qualname: str = None,
                      group=None, app_name: str = None) -> list[dict]:
    """生成某应用的内置文件工具定义（工具名 <prefix>__platform_*，prefix 为组限定前缀）。

    _meta.app 为 qualname（路由/鉴权用）；group/app_name 供展示。"""
    qn = qualname or prefix
    _app = app_name or qn
    return [
        {
            "name": f"{prefix}__platform_list_files",
            "description": (
                "列出该应用 resource/import-file/（人工上传的待导入文件）或 "
                "resource/export-file/（产出结果）目录下的文件（名称/大小/修改时间）。"
            ),
            "inputSchema": {
                "type": "object",
                "properties": {
                    "directory": {
                        "type": "string", "enum": list(ALLOWED_DIRS),
                        "description": "import-file=待导入；export-file=产出结果。默认 import-file",
                        "default": IMPORT_DIR,
                    }
                },
                "required": [],
            },
            "_meta": {"app": qn, "service": "platform_list_files", "group": group, "app_name": _app},
        },
        {
            "name": f"{prefix}__platform_read_file",
            "description": (
                "读取该应用 resource 下的文件内容：文本(csv/tsv/txt/json/md 等)按 UTF-8 读取；"
                "xlsx 转首个工作表为文本表格（需 openpyxl）；过大只返回开头并标记 truncated。"
                "用于解析导入文件。"
            ),
            "inputSchema": {
                "type": "object",
                "properties": {
                    "directory": {"type": "string", "enum": list(ALLOWED_DIRS),
                                  "description": "import-file 或 export-file"},
                    "file_name": {"type": "string", "description": "文件名（可先用 platform_list_files 查看）"},
                },
                "required": ["directory", "file_name"],
            },
            "_meta": {"app": qn, "service": "platform_read_file", "group": group, "app_name": _app},
        },
        {
            "name": f"{prefix}__platform_write_file",
            "description": (
                "向该应用 resource/export-file/ 写入文本文件（如导入结果、分析报告、整理后的数据）。"
                "import-file 为人工上传区，不可写。append=true 追加，默认覆盖。"
            ),
            "inputSchema": {
                "type": "object",
                "properties": {
                    "file_name": {"type": "string", "description": "文件名（含扩展名）"},
                    "content": {"type": "string", "description": "文本内容"},
                    "append": {"type": "boolean", "description": "true=追加；默认覆盖", "default": False},
                },
                "required": ["file_name", "content"],
            },
            "_meta": {"app": qn, "service": "platform_write_file", "group": group, "app_name": _app},
        },
    ]
